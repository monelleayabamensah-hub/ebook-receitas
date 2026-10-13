"""Monta a versão editável para o Canva (.pptx) a partir das camadas do canva-export.js.

Cada página vira um slide A4 com: fundo (sem texto), traçados, ilustrações
soltas e caixas de texto editáveis com as fontes do ebook.

Uso:  node canva-export.js <pasta>  &&  python3 canva_pptx.py <pasta>
"""
import json
import re
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

PX = 9525  # EMU por pixel CSS (96 dpi)
W, H = 794, 1123  # A4 em px

FONTS = {  # (família CSS, peso) -> nome da fonte instalada / enviada ao Canva
    ("Barlow Condensed", "800"): "Barlow Condensed ExtraBold",
    ("Barlow Condensed", "600"): "Barlow Condensed SemiBold",
    ("Short Stack", "400"): "Short Stack",
}
ALIGN = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "end": PP_ALIGN.RIGHT}


def font_name(run):
    family = run["family"].split(",")[0].strip().strip('"')
    return FONTS.get((family, run["weight"]), family)


def rgb(css):
    r, g, b = (int(v) for v in re.findall(r"\d+", css)[:3])
    return RGBColor(r, g, b)


def add_text(slide, block):
    b = block["box"]
    # um pouco de folga na largura: o Canva/PowerPoint quebra linha de forma levemente diferente do navegador
    tb = slide.shapes.add_textbox(Emu(int(b["x"] * PX)), Emu(int(b["y"] * PX)),
                                  Emu(int(b["w"] * 1.04 * PX + 2 * PX)), Emu(int(b["h"] * PX)))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    para = tf.paragraphs[0]

    def style_para(p):
        p.line_spacing = Pt(block["lineHeight"] * 0.75)
        if block["align"] in ALIGN:
            p.alignment = ALIGN[block["align"]]

    style_para(para)
    for rn in block["runs"]:
        if rn.get("br"):
            para = tf.add_paragraph()
            style_para(para)
            continue
        r = para.add_run()
        r.text = rn["t"]
        f = r.font
        f.name = font_name(rn)
        f.size = Pt(round(rn["size"] * 0.75 * 2) / 2)
        f.color.rgb = rgb(rn["color"])
        if rn["ls"] not in ("normal", "0px"):
            r._r.get_or_add_rPr().set("spc", str(int(float(rn["ls"][:-2]) * 0.75 * 100)))


def add_picture(slide, path, x, y, w, h, name):
    pic = slide.shapes.add_picture(str(path), Emu(int(x * PX)), Emu(int(y * PX)), Emu(int(w * PX)), Emu(int(h * PX)))
    pic.name = name
    return pic


def main(src):
    src = Path(src)
    pages = json.loads((src / "layout.json").read_text())
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W * PX), Emu(H * PX)
    blank = prs.slide_layouts[6]
    for i, pg in enumerate(pages, start=1):
        s = prs.slides.add_slide(blank)
        bg = add_picture(s, src / f"p{i}-fundo.png", 0, 0, W, H, "Fundo")
        add_picture(s, src / f"p{i}-tracados.png", 0, 0, W, H, "Traçados")
        for k, a in enumerate(pg["art"], start=1):
            bx = a["box"]
            add_picture(s, src / f"p{i}-ilustracao{k}.png", bx["x"], bx["y"], bx["w"], bx["h"], f"Ilustração {k}")
        for block in pg["blocks"]:
            add_text(s, block)
        # trava o fundo para não sair do lugar ao editar
        nv = bg._element.find(qn("p:nvPicPr")).find(qn("p:cNvPicPr"))
        locks = nv.find(qn("a:picLocks"))
        if locks is None:
            locks = etree.SubElement(nv, qn("a:picLocks"))
        locks.set("noMove", "1")
        locks.set("noResize", "1")
    out = Path(__file__).with_name("canva") / "cozinha-mamede-canva.pptx"
    prs.save(out)
    print(f"{len(pages)} slides -> {out}")


if __name__ == "__main__":
    main(sys.argv[1])
